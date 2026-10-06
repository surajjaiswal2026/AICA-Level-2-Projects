"""Excel registers (openpyxl): bill register and outstanding register."""
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1556B0")
TOTAL_FILL = PatternFill("solid", fgColor="EAF1FB")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="B8C7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0.00"


def _to_date(iso):
    try:
        return datetime.strptime(iso, "%Y-%m-%d").date()
    except Exception:
        return iso or ""


def _sheet(wb, title, heading, columns, rows, money_cols, total_cols, first=False):
    """columns: list of (header, width). Adds title, header, rows and a total row."""
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.cell(row=1, column=1, value=heading).font = Font(bold=True, size=14, color="1556B0")
    ws.cell(row=2, column=1, value=f"Generated on {datetime.now():%d-%m-%Y %H:%M}").font = Font(italic=True, color="555555")
    head_row = 4
    for col, (name, width) in enumerate(columns, 1):
        cell = ws.cell(row=head_row, column=col, value=name)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(rows, head_row + 1):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            if c in money_cols:
                cell.number_format = MONEY
            elif isinstance(value, date):
                cell.number_format = "DD-MM-YYYY"
    if rows and total_cols:
        total_row = head_row + len(rows) + 1
        ws.cell(row=total_row, column=1, value="TOTAL")
        for c in range(1, len(columns) + 1):
            cell = ws.cell(row=total_row, column=c)
            cell.fill, cell.font, cell.border = TOTAL_FILL, Font(bold=True), BORDER
            if c in total_cols:
                letter = get_column_letter(c)
                cell.value = f"=SUM({letter}{head_row + 1}:{letter}{total_row - 1})"
                cell.number_format = MONEY
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    return ws


def _invoice_row(inv):
    cust = inv.get("customer", {})
    return [inv.get("invoice_no", ""), _to_date(inv.get("date", "")),
            cust.get("name", ""), cust.get("mobile", ""), cust.get("gstin", ""),
            cust.get("address", "").replace("\n", ", "),
            inv.get("subtotal", 0), inv.get("cgst", 0), inv.get("sgst", 0),
            inv.get("igst", 0), inv.get("total_gst", 0), inv.get("round_off", 0),
            inv.get("grand_total", 0), inv.get("received_amount", 0),
            inv.get("outstanding_amount", 0), inv.get("payment_status", ""),
            _to_date(inv.get("last_receipt_date", ""))]


INVOICE_COLS = [("Invoice No", 14), ("Invoice Date", 13), ("Customer Name", 28),
                ("Mobile", 14), ("Customer GSTIN", 19), ("Address", 34),
                ("Taxable Value", 15), ("CGST", 12), ("SGST", 12), ("IGST", 12),
                ("Total GST", 13), ("Round Off", 11), ("Grand Total", 15),
                ("Received", 14), ("Outstanding", 14), ("Status", 13),
                ("Last Receipt Date", 16)]
INVOICE_MONEY = set(range(7, 16))


def _item_rows(invoices):
    rows = []
    for inv in invoices:
        for it in inv.get("items", []):
            rows.append([inv.get("invoice_no", ""), _to_date(inv.get("date", "")),
                         inv.get("customer", {}).get("name", ""), it.get("name", ""),
                         it.get("hsn", ""), it.get("qty", 0), it.get("rate", 0),
                         it.get("amount", 0), it.get("gst_pct", 0),
                         it.get("gst_amount", 0), it.get("total", 0)])
    return rows


def _receipt_rows(invoices):
    rows = []
    for inv in invoices:
        for rc in inv.get("receipts", []):
            rows.append([inv.get("invoice_no", ""), inv.get("customer", {}).get("name", ""),
                         _to_date(rc.get("date", "")), rc.get("amount", 0),
                         rc.get("mode", ""), rc.get("note", "")])
    return rows


ITEM_COLS = [("Invoice No", 14), ("Invoice Date", 13), ("Customer Name", 28),
             ("Item Name", 32), ("HSN/SAC", 12), ("Qty", 10), ("Rate", 13),
             ("Amount", 15), ("GST %", 9), ("GST Amount", 14), ("Total", 15)]
RECEIPT_COLS = [("Invoice No", 14), ("Customer Name", 28), ("Receipt Date", 14),
                ("Amount", 15), ("Mode", 14), ("Note / Reference", 34)]


def export_bill_register(invoices, path, seller_name=""):
    invoices = sorted(invoices, key=lambda i: (i.get("date", ""), str(i.get("invoice_no", ""))))
    wb = Workbook()
    title = f"{seller_name} - Bill Register".strip(" -")
    _sheet(wb, "Bill Register", title, INVOICE_COLS, [_invoice_row(i) for i in invoices],
           INVOICE_MONEY, INVOICE_MONEY, first=True)
    _sheet(wb, "Item Wise", f"{seller_name} - Item Wise Details".strip(" -"), ITEM_COLS,
           _item_rows(invoices), {7, 8, 10, 11}, {8, 10, 11})
    _sheet(wb, "Receipts", f"{seller_name} - Receipts".strip(" -"), RECEIPT_COLS,
           _receipt_rows(invoices), {4}, {4})
    wb.save(str(path))
    return str(path)


def export_outstanding_register(invoices, path, seller_name=""):
    invoices = [i for i in invoices if i.get("outstanding_amount", 0) > 0.004]
    invoices.sort(key=lambda i: (i.get("date", ""), str(i.get("invoice_no", ""))))
    today = date.today()
    cols = INVOICE_COLS + [("Days Outstanding", 12)]
    rows = []
    for inv in invoices:
        d = _to_date(inv.get("date", ""))
        rows.append(_invoice_row(inv) + [(today - d).days if isinstance(d, date) else ""])
    wb = Workbook()
    _sheet(wb, "Outstanding Register", f"{seller_name} - Outstanding Bill Register".strip(" -"),
           cols, rows, INVOICE_MONEY, INVOICE_MONEY, first=True)
    _sheet(wb, "Item Wise", "Item Wise Details of Outstanding Invoices", ITEM_COLS,
           _item_rows(invoices), {7, 8, 10, 11}, {8, 10, 11})
    _sheet(wb, "Part Receipts", "Part Receipts against Outstanding Invoices", RECEIPT_COLS,
           _receipt_rows(invoices), {4}, {4})
    wb.save(str(path))
    return str(path)
